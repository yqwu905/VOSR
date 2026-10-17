"""Synthetic Chinese slides: real template decks refilled with real Chinese sentences.

Chinese decks are scarce, while template decks (ypppt, the 四川大学 templates on Zenodo) have
real layouts, themes and fonts but 12-56% placeholder text (单击此处添加标题, 请替换文字内容,
lorem ipsum, English filler). ``refill`` keeps every shape and the formatting of the first
non-empty run of a paragraph and swaps the text of such paragraphs for corpus text of about the
same visual width (a CJK character is about twice as wide as a Latin one); lines naming
template sites (1ppt.com, ypppt.com, 51pptmoban, ...) are blanked. Slide-master and layout
placeholders are left alone (their prompt text never renders); other text on them is treated
like slide text. The corpus should be text whose terms allow it (e.g. Art. 5 government
documents); the layouts keep the template's own terms.
"""

import io
import random
import re
from bisect import bisect_left, bisect_right

from .sniff import cjk_ratio

PLACEHOLDER = re.compile(
    r"单击此处|点击此处|单击添加|点击添加|此处添加|此处输入|请在此|请输入|请替换|输入标题|输入小标题|添加标题|副标题|标题文本|"
    r"文本标题|标题文字|添加文字|添加文本|添加内容|输入文字|输入内容|输入[你您]|[你您]的内容|[你您](?:需要|想要)的(?:文|内容)|"
    r"内容打在这里|替换文字|替换内容|在这里输入|这里输入|[你您]的标题|此处键入|在此键入|在此录入|编辑母版|文字大小|"
    r"一键替换|替换和编辑|原字体|使用开心|根据您的(?:个人)?需要|"
    r"lorem|ipsum|dolor sit|your (?:text|title|content|logo)|logo here|add (?:your )?(?:title|text|content)|"
    r"click (?:here )?to (?:add|edit)|enter (?:your )?(?:title|text)|(?:text|title) here|XXX+|"
    r"(?<!\d)20[\dX]X(?![A-Za-z\d])", re.I)
WATERMARK = re.compile(r"1ppt\.(?:com|cn)|ypppt\.com|51pptmoban|第一PPT|优品PPT|pptstore|docer\.com|tukuppt|ibaotu|588ku|"
                       r"51miz|pptbz|islide", re.I)
URL = re.compile(r"https?://|www\.|\.com\b|\.cn\b", re.I)
CJK = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")
WIDE = re.compile(r"[\u1100-\u115f\u2e80-\ua4cf\uac00-\ud7a3\uf900-\ufaff\ufe30-\ufe4f\uff00-\uff60\uffe0-\uffe6]")
CLAUSE = re.compile(r"[，、；：,;:]")


class Corpus:
    """Sentences and their clauses (for short titles), drawn by length."""

    def __init__(self, sentences, rng):
        pieces = set()
        for sentence in sentences:
            sentence = re.sub(r"\s+", "", sentence)
            if len(sentence) < 2 or URL.search(sentence) or WATERMARK.search(sentence) or not CJK.search(sentence):
                continue
            pieces.add(sentence)
            pieces.update(p for p in CLAUSE.split(sentence) if len(p) >= 2 and CJK.search(p))
        if not pieces:
            raise ValueError("refill needs Chinese sentences without URLs or template-site names")
        self.items = sorted(pieces, key=lambda s: (len(s), s))
        self.lengths = [len(s) for s in self.items]
        self.rng = rng

    def take(self, n):
        """Text of about n characters: one piece within ±20% when possible, else pieces joined and cut."""
        low, high = bisect_left(self.lengths, round(n * 0.8)), bisect_right(self.lengths, max(n, round(n * 1.2)))
        if high > low:
            return self.items[self.rng.randrange(low, high)]
        text = ""
        while len(text) < n:
            text += self.rng.choice(self.items)
        return text[:n]


def width(text):
    """Visual width in CJK character units."""
    return sum(1.0 if WIDE.match(c) else 0.5 for c in text)


def latin_filler(text, min_chars=20):
    """English or lorem-ipsum body text: no CJK, mostly letters, a few words."""
    chars = [c for c in text if not c.isspace()]
    if len(chars) < min_chars or CJK.search(text) or len(text.split()) < 3:
        return False
    return sum(c.isascii() and c.isalpha() for c in chars) >= 0.6 * len(chars)


def _text_frames(shapes, skip_placeholders=False):
    from pptx.shapes.group import GroupShape
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _text_frames(shape.shapes, skip_placeholders)
        elif skip_placeholders and shape.is_placeholder:
            continue
        elif getattr(shape, "has_text_frame", False):
            yield shape.text_frame
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def refill(pptx_bytes, sentences, seed=0, latin=True, min_latin=20):
    """Return (new pptx bytes, stats) with placeholder and filler paragraphs refilled.

    ``sentences`` is an iterable of Chinese sentences; ``latin`` also refills English/lorem
    filler of at least ``min_latin`` non-space characters. Stats count slides, paragraphs,
    replaced placeholders, replaced Latin filler, blanked watermark lines and the CJK share
    of all slide text before and after."""
    from pptx import Presentation
    deck = Presentation(io.BytesIO(pptx_bytes))
    corpus = Corpus(sentences, random.Random(seed))
    stats = dict(slides=len(deck.slides), paragraphs=0, placeholder=0, latin=0, watermark=0)
    frames = [frame for slide in deck.slides for frame in _text_frames(slide.shapes)]
    for master in deck.slide_masters:
        frames += _text_frames(master.shapes, skip_placeholders=True)
        for layout in master.slide_layouts:
            frames += _text_frames(layout.shapes, skip_placeholders=True)
    before, after = [], []
    for frame in frames:
        for paragraph in frame.paragraphs:
            runs = paragraph.runs
            text = "".join(run.text for run in runs)
            if not text.strip():
                continue
            stats["paragraphs"] += 1
            before.append(text)
            if WATERMARK.search(text):
                kind, new = "watermark", ""
            elif PLACEHOLDER.search(text):
                kind, new = "placeholder", corpus.take(max(2, round(width(text.strip()))))
            elif latin and latin_filler(text, min_latin):
                kind, new = "latin", corpus.take(max(2, round(width(text.strip()))))
            else:
                after.append(text)
                continue
            stats[kind] += 1
            keep = next(run for run in runs if run.text.strip())
            for run in runs:
                run.text = new if run is keep else ""
            after.append(new)
    stats.update(replaced=stats["placeholder"] + stats["latin"], cjk_before=round(cjk_ratio("".join(before)), 3),
                 cjk_after=round(cjk_ratio("".join(after)), 3))
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue(), stats
