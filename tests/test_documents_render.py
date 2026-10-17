"""Document conversion and rendering: font aliases, macro refusal, page classes, DPI targeting, boxes."""

import json
import os
import shutil
import subprocess
import zipfile

import numpy as np
import pytest

pymupdf = pytest.importorskip("pymupdf")

from data_pipeline.documents.convert import ConversionError, fonts_conf, has_macros, prepare_profile, to_pdf
from data_pipeline.documents.fetch import Store
from data_pipeline.documents.render import classify, font_name, page_lines, render_store

ZH = "关于进一步优化营商环境的若干措施"


def installed(family):
    out = subprocess.run(["fc-list", ":", "family"], capture_output=True, text=True).stdout if shutil.which("fc-list") else ""
    return family in out


@pytest.mark.skipif(not installed("Noto Serif CJK SC") or not installed("AR PL UKai CN"), reason="CJK fonts missing")
def test_font_aliases_keep_song_kai_and_hei_styles(tmp_path):
    env = dict(os.environ, FONTCONFIG_FILE=str(fonts_conf(tmp_path)))
    match = lambda family: subprocess.run(["fc-match", family, "family"], capture_output=True, text=True,
                                          env=env).stdout.strip()
    assert match("微软雅黑") == "Noto Sans CJK SC"
    assert match("仿宋_GB2312") == "Noto Serif CJK SC"
    assert match("楷体") == "AR PL UKai CN"
    assert match("汉仪中黑简") == "Noto Sans CJK SC"  # unknown family, matched by name fragment
    assert match("方正小标宋_GBK") == "Noto Serif CJK SC"
    assert match("Arial") != "Noto Sans CJK SC"  # Latin fonts are untouched


def test_macro_documents_are_refused(tmp_path):
    path = tmp_path / "deck.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("ppt/presentation.xml", "x")
        archive.writestr("ppt/vbaProject.bin", "x")
    assert has_macros(path)
    with pytest.raises(ConversionError, match="macros"):
        to_pdf(path, tmp_path / "out", tmp_path / "profile", tmp_path / "fonts.conf")


def test_font_name_cleanup():
    assert font_name("BAAAAA+NotoSerifCJKsc-Regular") == "NotoSerifCJKsc-Regular"
    assert font_name("宋体".encode("gbk").decode("latin-1")) == "宋体"


def pdf_with(pages):
    document = pymupdf.open()
    for build in pages:
        page = document.new_page(width=595, height=842)
        build(page)
    return document


def text_page(page, size=12, text=ZH):
    for i in range(12):
        page.insert_text((60, 80 + i * 2.2 * size), text, fontsize=size, fontname="china-s")


def test_page_classes():
    def scanned(page):
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 280), False)
        pix.clear_with(200)
        page.insert_image(page.rect, pixmap=pix)

    def vector(page):
        for i in range(20):
            page.draw_rect(pymupdf.Rect(60, 60 + i * 30, 400, 75 + i * 30), color=(0, 0, 0), fill=(0, 0, 0))

    document = pdf_with([text_page, scanned, vector, lambda page: None])
    found = []
    for page in document:
        thumb = page.get_pixmap(matrix=pymupdf.Matrix(0.5, 0.5), alpha=False)
        gray = np.frombuffer(thumb.samples, np.uint8).reshape(thumb.height, thumb.width, 3).mean(axis=2)
        found.append(classify(page, page_lines(page), float((gray < 160).mean()))[0])
    assert found == ["text", "scanned", "vector", "blank"]
    blank = document[3]
    private_use = [{"text": "\ue001\ue002\ue003 第一章 总则 \ue004\ue005\ue006", "font": "SimSun"}]
    founder = [{"text": "!\"#$%&'()*+,-./:;<=>? ;:!", "font": "FzBookMaker0DlFont005368"}]
    assert classify(blank, private_use, 0.05)[0] == "garbled" == classify(blank, founder, 0.05)[0]


def store_with_pdf(tmp_path, document, source="test"):
    store = Store(tmp_path / "store")
    digest, kind, path = store.put(document.tobytes())
    store.record({"key": "k", "source": source, "status": "ok", "sha256": digest, "type": kind, "path": str(path),
                  "terms": "official", "title": "t"})
    return store


def test_render_targets_text_height_and_boxes_cover_ink(tmp_path):
    def rotated(page):
        text_page(page, 10)
        page.set_rotation(90)

    def placeholder(page):
        page.insert_text((60, 80), "单击此处添加标题", fontsize=24, fontname="china-s")
        text_page(page, 12)

    document = pdf_with([lambda p: text_page(p, 12), lambda p: text_page(p, 9), rotated, placeholder,
                         lambda p: text_page(p, 12)])
    store_with_pdf(tmp_path, document)
    summary = render_store(tmp_path / "store", tmp_path / "out", text_height="40", workers=1, log=None)
    rows = [json.loads(line) for line in (tmp_path / "out" / "clean.jsonl").read_text().splitlines()]
    pages = {json.loads(line)["document"]["page"]: json.loads(line)
             for line in (tmp_path / "out" / "pages.jsonl").read_text().splitlines()}
    assert pages[3]["status"] == "dropped_template" and pages[4]["status"] == "dropped_duplicate"
    assert summary["pages"]["kept"] == 3 == len(rows)
    for row in rows:
        assert abs(row["median_text_px"] - 40) < 2
        from PIL import Image
        image = np.asarray(Image.open(row["image_path"]).convert("L"))
        assert list(image.shape[::-1]) == row["size"]
        for item in row["annotations"]:
            xs, ys = zip(*item["polygon"])
            box = image[int(min(ys)):int(max(ys)) + 1, int(min(xs)):int(max(xs)) + 1]
            assert item["trusted"] and item["text"] == ZH and box.size and (box < 128).mean() > 0.05


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice missing")
def test_word_document_is_converted_and_labelled(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn
    document = docx.Document()
    for _ in range(6):
        run = document.add_paragraph().add_run(ZH)
        run.font.name = "仿宋_GB2312"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋_GB2312")
    path = tmp_path / "notice.docx"
    document.save(path)
    store = Store(tmp_path / "store")
    digest, kind, stored = store.put(path.read_bytes())
    store.record({"key": "k", "source": "t", "status": "ok", "sha256": digest, "type": kind, "path": str(stored),
                  "terms": "official"})
    render_store(tmp_path / "store", tmp_path / "out", workers=1, log=None)
    row = json.loads((tmp_path / "out" / "clean.jsonl").read_text().splitlines()[0])
    assert row["class"] == "text" and row["annotations"][0]["text"] == ZH
    assert "NotoSerifCJK" in row["annotations"][0]["font"]  # 仿宋 rendered in a Song-style face
