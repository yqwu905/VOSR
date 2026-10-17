"""Slide sources (offline): WebPlus course pages, VALSE, Zenodo CJK decks, ypppt; archive
members; template refill."""

import io
import json
import struct
import zipfile
import zlib
from types import SimpleNamespace

import pytest

from data_pipeline.documents import catalog, slides
from data_pipeline.documents.fetch import FetchError
from data_pipeline.documents.unpack import members


class FakeFetcher:
    """Canned pages by URL; a value may be (final_url, body) to simulate a redirect."""

    def __init__(self, pages):
        self.pages, self.calls = pages, []

    def get(self, url, headers=None, **kwargs):
        self.calls.append((url, dict(headers or {})))
        page = self.pages.get(url)
        if page is None:
            raise FetchError("http_404", url)
        final, body = page if isinstance(page, tuple) else (url, page)
        return SimpleNamespace(url=final, status_code=200, headers={}), body.encode() if isinstance(body, str) else body

    def json(self, url, **kwargs):
        return json.loads(self.get(url, **kwargs)[1])

    def urls(self):
        return [url for url, _ in self.calls]


# --- WebPlus ---

U = "https://www.u.edu.cn"
SEED = U + "/szb/2020/0701/c9141a168693/page.htm"
FILES = "/_upload/article/files/0c/38/6cd63d57424aab01c67b4b37211e/"

SEED_PAGE = f"""<html><head><title>《毛泽东思想概论》PPT</title></head><body>
<a href="/szb/jxkj/list.htm"><span class="column-name">教学课件</span></a>
<a href="/szb/xwdt/list.htm">新闻动态</a><a href="/zs/list.htm">招生</a><a href="javascript:;">x</a>
<a href="http://www.u.edu.cn/szb/jxdt/list.htm">教学动态</a>
<a href="https://other.edu.cn/szb/jxkj/list.htm">兄弟院校课件</a>
<a href="/szb/2020/0701/c9141a168694/page.htm" title="学生名单公示">学生名单公示</a>
<p><a href="{FILES}ffa12f5c.ppt" sudyfile-attr="{{'title':'01第一章&nbsp;&nbsp;毛泽东思想及其历史地位.ppt'}}">01第一章</a></p>
<p><a href="{FILES}ffa12f5c.ppt" sudyfile-attr="{{'title':'01第一章.ppt'}}">again</a></p>
<p><a href="{FILES}d107d19a.doc" sudyfile-attr="{{'title':'习题解答.doc'}}">习题解答.doc</a></p>
<p><a href="{FILES}aaaa.pdf" sudyfile-attr="{{'title':'2020级学生名单.pdf'}}">名单</a></p>
<div class="wp_pdf_player" pdfsrc="{FILES}bbbb.pdf" sudyfile-attr="{{'title':'第二章 新民主主义革命理论.pdf'}}"></div>
<a href="http://www.u.edu.cn{FILES}cccc.PPTX">第三章 社会主义改造理论.pptx</a>
<a href="/images/logo.ppt">not an upload</a>
</body></html>"""

LISTING = """<html><title>教学课件</title>
<a href='/szb/2020/0630/c9141a168527/page.htm' title='《思想道德修养与法律基础》PPT'>思修</a>
<a href='/szb/2020/0629/c9141a168500/page.htm' title='学院新闻'>学院新闻</a>
<a href='/_upload/article/files/c5/be/9033804c/ee90202a.pptx' title='《中国近现代史纲要》PPT'>纲要</a>
<em class="curr_page">1</em>/<em class="all_pages">2</em></html>"""

ARTICLE = f"""<title>《思想道德修养与法律基础》PPT</title>
<a href="{FILES}eeee.pptx" sudyfile-attr="{{'title':'绪论.pptx'}}">绪论</a>"""


def test_webplus_seed_page_attachments_titles_filters_and_dedupe():
    fake = FakeFetcher({SEED: SEED_PAGE})
    found = list(slides.webplus(fake, {"pages": [SEED]}))
    assert [c["key"] for c in found] == [f"www.u.edu.cn{FILES}ffa12f5c.ppt", f"www.u.edu.cn{FILES}bbbb.pdf",
                                         f"www.u.edu.cn{FILES}cccc.PPTX"]
    first, player, plain = found
    assert first["title"] == "01第一章 毛泽东思想及其历史地位" and first["type_hint"] == "ppt"
    assert first["headers"] == {"Referer": SEED} and first["meta"]["published"] == "2020-07-01"
    assert first["meta"]["page_title"] == "《毛泽东思想概论》PPT" and first["license"] == slides.WEBPLUS_LICENSE
    assert player["title"] == "第二章 新民主主义革命理论" and player["type_hint"] == "pdf"
    assert plain["url"] == f"https://www.u.edu.cn{FILES}cccc.PPTX" and plain["type_hint"] == "pptx"
    assert plain["title"] == "第三章 社会主义改造理论"
    assert fake.urls() == [SEED]  # max_pages defaults to the seed alone
    pdf_only = list(slides.webplus(FakeFetcher({SEED: SEED_PAGE}), {"pages": SEED, "extensions": "pdf"}))
    assert [c["type_hint"] for c in pdf_only] == ["pdf"]


def test_webplus_crawl_prefers_courseware_listings_follows_pages_and_stays_in_scope():
    pages = {SEED: SEED_PAGE, U + "/szb/jxkj/list.htm": LISTING, U + "/szb/jxkj/list2.htm": "<title>教学课件</title>",
             U + "/szb/2020/0630/c9141a168527/page.htm": ARTICLE, U + "/szb/xwdt/list.htm": "<title>新闻</title>",
             U + "/szb/jxdt/list.htm": "<title>教学动态</title>"}
    fake = FakeFetcher(pages)
    found = list(slides.webplus(fake, {"pages": [SEED], "max_pages": 5}))
    # Links naming slides first (pagination ranks like its listing), then course material, never plain news here.
    assert fake.urls() == [SEED, U + "/szb/jxkj/list.htm", U + "/szb/2020/0630/c9141a168527/page.htm",
                           U + "/szb/jxkj/list2.htm", U + "/szb/jxdt/list.htm"]
    keys = [c["key"] for c in found]
    assert "www.u.edu.cn/_upload/article/files/c5/be/9033804c/ee90202a.pptx" in keys
    listed = found[keys.index("www.u.edu.cn/_upload/article/files/c5/be/9033804c/ee90202a.pptx")]
    assert listed["title"] == "《中国近现代史纲要》PPT" and listed["meta"]["page"] == U + "/szb/jxkj/list.htm"
    assert keys[-1] == f"www.u.edu.cn{FILES}eeee.pptx" and len(keys) == len(set(keys)) == 5
    # Out of prefix (/zs/), other hosts and personal-data pages are never fetched.
    assert not any("/zs/" in u or "other.edu.cn" in u or "c9141a168694" in u for u in fake.urls())
    wide = FakeFetcher(dict(pages, **{U + "/zs/list.htm": "<title>招生</title>"}))
    list(slides.webplus(wide, {"pages": [SEED], "max_pages": 10, "prefix": "/"}))
    assert U + "/zs/list.htm" in wide.urls() and U + "/szb/xwdt/list.htm" in wide.urls()


def test_webplus_unreachable_seed_and_item_cap():
    fake = FakeFetcher({SEED: SEED_PAGE})
    assert list(slides.webplus(fake, {"pages": ["http://down.edu.cn/x/page.htm"]})) == []
    assert len(list(slides.webplus(fake, {"pages": [SEED], "max_items": 2}))) == 2
    assert slides._subsite("/2019/0729/c15638a254307/page.htm") == "/"
    assert slides._subsite("/szb/jxkj/list.htm") == "/szb/" and slides._subsite("/15638/list.htm") == "/"


# --- VALSE ---

VROOT = slides.VALSE_ROOT
VDIR = "https://valser.org/webinar/slide/index.php/Home/Index/index/dir/"
VFILES = "http://valser.org/webinar/slide/slides/"


def vdir_link(name):
    return f'<td><a href="/webinar/slide/index.php/Home/Index/index/dir/{name}.html">{name}</a></td>'


def vfile_link(path):
    preview = "/webinar/slide/index.php/Home/Index/index/file/%2Fvalser%2F/name/x.pdf.html"
    return f'<a href="{preview}">preview</a><a target="_blank"  href="{VFILES}{path}">{path}</a>'


VALSE_PAGES = {
    VROOT: "".join([vdir_link("20260107"), vdir_link("20251217"), vdir_link("%E7%9F%AD%E6%95%99%E7%A8%8B06"),
                    vdir_link("20190703"), vdir_link("20260107"), vfile_link("/software_cracking.pdf")]),
    VDIR + "20260107.html": vfile_link("20260107/可信视觉鉴伪大模型-张健-20260107-v2.pdf")
    + vfile_link("20260107/talk.pptx") + vfile_link("20260107/Как_сбежать.pdf"),
    VDIR + "20251217.html": vfile_link("20251217/200520_VALSE_Deep Lesion &amp; mining.pdf"),
    VDIR + "%E7%9F%AD%E6%95%99%E7%A8%8B06.html": vfile_link("短教程06/视觉与语言.pdf") + vdir_link("20251217"),
    VDIR + "20190703.html": vfile_link("20190703/relation_networks.pdf"),
}


def test_valse_walks_directories_quotes_names_and_filters():
    fake = FakeFetcher(VALSE_PAGES)
    found = list(slides.valse(fake, {}))
    assert [c["key"] for c in found] == ["valse:20260107/可信视觉鉴伪大模型-张健-20260107-v2.pdf",
                                         "valse:20251217/200520_VALSE_Deep Lesion & mining.pdf",
                                         "valse:短教程06/视觉与语言.pdf", "valse:20190703/relation_networks.pdf"]
    assert found[0]["url"] == ("https://valser.org/webinar/slide/slides/20260107/%E5%8F%AF%E4%BF%A1%E8%A7%86%E8%A7%89"
                               "%E9%89%B4%E4%BC%AA%E5%A4%A7%E6%A8%A1%E5%9E%8B-%E5%BC%A0%E5%81%A5-20260107-v2.pdf")
    assert found[1]["url"].endswith("/20251217/200520_VALSE_Deep%20Lesion%20%26%20mining.pdf")
    assert found[0]["title"] == "可信视觉鉴伪大模型-张健-20260107-v2" and found[0]["meta"]["date"] == "2026-01-07"
    assert found[2]["meta"]["date"] is None  # root files and Cyrillic junk uploads are skipped
    assert fake.urls().count(VDIR + "20251217.html") == 1  # a directory linked twice is read once

    recent = FakeFetcher(VALSE_PAGES)
    keys = [c["key"] for c in slides.valse(recent, {"years": "2025-2026", "undated": False})]
    assert keys == ["valse:20260107/可信视觉鉴伪大模型-张健-20260107-v2.pdf",
                    "valse:20251217/200520_VALSE_Deep Lesion & mining.pdf"]
    assert VDIR + "20190703.html" not in recent.urls()
    chinese = [c["key"] for c in slides.valse(FakeFetcher(VALSE_PAGES), {"cjk_names": True, "extensions": "pdf,pptx"})]
    assert chinese == ["valse:20260107/可信视觉鉴伪大模型-张健-20260107-v2.pdf", "valse:短教程06/视觉与语言.pdf"]
    capped = FakeFetcher(VALSE_PAGES)
    assert len(list(slides.valse(capped, {"max_dirs": 1}))) == 1 and len(capped.urls()) == 2
    assert list(slides.valse(FakeFetcher({}), {})) == []


# --- Zenodo ---

def zenodo_parquet(rows):
    import pyarrow as pa
    import pyarrow.parquet as pq
    columns = ["filename", "size", "url", "license", "title", "created", "doi", "checksum"]
    table = pa.table({c: [str(row.get(c, "")) if c != "size" else int(row.get(c, 0)) for row in rows] for c in columns})
    buffer = io.BytesIO()
    pq.write_table(table, buffer)
    return buffer.getvalue()


def zrow(record, filename, title, license="cc-by-4.0", size=1000):
    return dict(filename=filename, size=size, url=f"https://zenodo.org/api/records/{record}/files/{filename}/content",
                license=license, title=title, created="2021-01-29T06:58:20+00:00", doi=f"10.5281/zenodo.{record}",
                checksum="md5:00")


ZROWS = [zrow(1207215, "UF DSI Symposium Talk.pptx", "Data Science"),
         zrow(4478221, "年度计划.pptx", "xxx"),
         zrow(5089075, "王_2017_学習者が犯した誤用.pptx", "学習者が犯した誤用の原因は何か"),
         zrow(4289865, "talk.pptx", "听声识物与辨物知声", license="cc-by-nc-4.0"),
         zrow(7791547, "蓝花楹.pptx", "四川大学PPT模板", size=223233777),
         zrow(7791547, "荷花粉.pptx", "四川大学PPT模板", license="cc-zero", size=11689810),
         zrow(7791547, "荷花粉.pptx", "四川大学PPT模板", license="cc-zero", size=11689810)]

EXPORT = {"created": "2023-06-20T12:33:50+00:00", "pids": {"doi": {"identifier": "10.5281/zenodo.8059899"}},
          "metadata": {"title": "Manage your bibliography with Zotero (Chinese) - 文献管理软件",
                       "rights": [{"id": "cc-by-4.0"}], "languages": [{"id": "cmn"}]},
          "files": {"entries": {"Zotero-文献管理软件.pdf": {"size": 3212678, "checksum": "md5:de"},
                                "cover.png": {"size": 10, "checksum": "md5:01"}}}}


def test_zenodo_cjk_filters_rows_by_script_licence_and_size():
    fake = FakeFetcher({slides.ZENODO10K: zenodo_parquet(ZROWS)})
    found = list(slides.zenodo_cjk(fake, {}))
    assert [c["key"] for c in found] == ["zenodo:4478221/年度计划.pptx", "zenodo:7791547/荷花粉.pptx"]
    first = found[0]
    assert first["url"] == ("https://zenodo.org/records/4478221/files/%E5%B9%B4%E5%BA%A6%E8%AE%A1%E5%88%92.pptx"
                            "?download=1")
    assert first["title"] == "年度计划" and first["license"] == "cc-by-4.0" and first["type_hint"] == "pptx"
    assert first["meta"]["record"] == 4478221 and first["meta"]["doi"] == "10.5281/zenodo.4478221"
    assert found[1]["title"] == "四川大学PPT模板" and found[1]["license"] == "cc-zero"
    wider = [c["key"] for c in slides.zenodo_cjk(FakeFetcher({slides.ZENODO10K: zenodo_parquet(ZROWS)}),
                                                 {"kana": True, "max_mb": 300, "licenses": "cc-by-4.0"})]
    assert wider == ["zenodo:4478221/年度计划.pptx", "zenodo:5089075/王_2017_学習者が犯した誤用.pptx",
                     "zenodo:7791547/蓝花楹.pptx"]


def test_zenodo_cjk_local_metadata_and_extra_records(tmp_path):
    path = tmp_path / "meta.parquet"
    path.write_bytes(zenodo_parquet(ZROWS[:2]))
    closed = dict(EXPORT, metadata=dict(EXPORT["metadata"], rights=[{"id": "cc-by-nc-4.0"}]))
    fake = FakeFetcher({"https://zenodo.org/records/8059899/export/json": json.dumps(EXPORT),
                        "https://zenodo.org/records/1/export/json": json.dumps(closed)})
    found = list(slides.zenodo_cjk(fake, {"metadata": str(path), "records": [8059899, 1, 404]}))
    assert [c["key"] for c in found] == ["zenodo:4478221/年度计划.pptx", "zenodo:8059899/Zotero-文献管理软件.pdf"]
    extra = found[1]
    assert extra["type_hint"] == "pdf" and extra["license"] == "cc-by-4.0"
    assert extra["title"].endswith("文献管理软件") and extra["meta"]["languages"] == ["cmn"]
    assert extra["meta"]["doi"] == "10.5281/zenodo.8059899"
    assert all(not url.endswith(".parquet") for url in fake.urls())


# --- ypppt ---

Y = slides.YPPPT
INDEX = """<div class="b-menu clear">
<h3>一年级上册</h3><a href="/special/aoe/">aoe</a></div>
<div class="b-menu clear">
<h3>八年级上册</h3><a href="/special/beiying/">背影</a><a href="/article/2022/11838.html">得道多助，失道寡助</a>
</div>"""


def lesson(*posts):
    menu = '<div class="b-menu clear"><a href="/article/2022/11838.html">得道多助，失道寡助</a></div>'
    items = "".join(f'<li><a href="/article/2021/{aid}.html"class="img_preview" target="_blank"><img src="x.jpg"></a>'
                    f'<a href="/article/2021/{aid}.html" class="p-title" target="_blank">{title}</a><p>…</p></li>'
                    for aid, title in posts)
    return f'{menu}<ul class="posts clear">{items}</ul>'


def download(title, link, pages="23页"):
    anchor = f'<li><a href="{link}">下载地址1</a></li><li><a href="{link}">下载地址2</a></li>' if link else ""
    return (f"<h1>{title} - 下载页</h1><ul class=\"clear\"><li><i>分类：</i>PPT课件</li>"
            f"<li><i>比例：</i>宽屏16:9</li><li class=\"sh\"><i>页数：</i>{pages}</li><li><i>格式：</i>PPTX</li>"
            f"<li class=\"sh\"><i>大小：</i>337KB</li></ul><ul class=\"down clear\">{anchor}</ul>")


def d(aid):
    return f"{Y}/p/d.php?aid={aid}"


YPAGES = {Y + "/kejian/yuwen/": "\ufeff" + INDEX,  # served with a BOM
          Y + "/special/aoe/": lesson((100, "aoe课件1")),
          Y + "/special/beiying/": lesson((11738, "背影PPT课件8"), (6417, "背影PPT课件2"), (11738, "背影PPT课件8")),
          d(100): download("aoe课件1", ""),
          d(11738): download("背影PPT课件8", "https://down.ypppt.com/uploads/soft/211231/1-211231104309.zip"),
          d(6417): download("背影PPT课件2", "https://down.ypppt.com/uploads/soft/200814/1-200Q4191048.rar", "24页"),
          d(11838): download("得道多助，失道寡助PPT课件", "https://down.ypppt.com/uploads/soft/220101/1-1.pptx")}


def test_ypppt_courseware_lessons_download_pages_and_referers():
    fake = FakeFetcher(YPAGES)
    found = list(slides.ypppt(fake, {}))
    assert [c["key"] for c in found] == ["ypppt:11738", "ypppt:6417", "ypppt:11838"]  # aoe has no download link
    deck = found[0]
    assert deck["url"] == "https://down.ypppt.com/uploads/soft/211231/1-211231104309.zip" and deck["type_hint"] == "zip"
    assert deck["title"] == "背影PPT课件8" and deck["headers"] == {"Referer": d(11738)}
    assert deck["license"] == slides.YPPPT_LICENSE and found[1]["type_hint"] == "rar"
    assert deck["meta"] == dict(deck["meta"], aid=11738, grade="八年级上册", lesson="背影", pages=23, ratio="宽屏16:9",
                                format="PPTX", size="337KB", category="PPT课件", section="kejian",
                                article=Y + "/article/2021/11738.html", host="down.ypppt.com")
    assert found[2]["meta"]["lesson"] == "得道多助，失道寡助" and found[2]["type_hint"] == "pptx"
    assert (d(11738), {"Referer": Y + "/article/2021/11738.html"}) in fake.calls
    assert fake.urls().count(d(11738)) == 1

    eighth = FakeFetcher(YPAGES)
    assert [c["key"] for c in slides.ypppt(eighth, {"grades": "八年级", "max_items": 1})] == ["ypppt:11738"]
    assert Y + "/special/aoe/" not in eighth.urls()
    one = FakeFetcher(YPAGES)
    assert list(slides.ypppt(one, {"lessons": ["aoe"]})) == [] and d(100) in one.urls()
    first = FakeFetcher(YPAGES)
    assert list(slides.ypppt(first, {"max_lessons": 1})) == [] and Y + "/special/beiying/" not in first.urls()


def listing(*posts):
    items = "".join(f'<li><a href="/article/2019/{aid}.html" class="p-title" target="_blank">{title}</a>'
                    f'<div class="clear"><span class="left">动态 - 25页</span></div></li>' for aid, title in posts)
    return f'<ul class="posts clear">{items}</ul><div class="page-navi"><a href="list-2.html">2</a></div>'


def test_ypppt_template_listing_pages_until_nothing_new():
    pages = {Y + "/moban/": listing((5722, "古典文化精美中国风PPT模板"), (5677, "精美深蓝粒子科技风PPT模板")),
             Y + "/moban/list-2.html": listing((5677, "精美深蓝粒子科技风PPT模板"), (5621, "创意黑板彩色粉笔PPT模板")),
             Y + "/moban/list-3.html": listing((5621, "创意黑板彩色粉笔PPT模板")),
             d(5722): download("古典文化精美中国风PPT模板", "https://down.ypppt.com/uploads/soft/191218/1-1.zip"),
             d(5677): download("精美深蓝粒子科技风PPT模板", "https://down.ypppt.com/uploads/soft/190929/1-2.zip"),
             d(5621): download("创意黑板彩色粉笔PPT模板", "https://down.ypppt.com/uploads/soft/190621/1-3.rar")}
    fake = FakeFetcher(pages)
    found = list(slides.ypppt(fake, {"sections": "moban", "max_list_pages": 10}))
    assert [c["key"] for c in found] == ["ypppt:5722", "ypppt:5677", "ypppt:5621"]
    assert found[0]["meta"]["section"] == "/moban/" and found[0]["meta"]["listing"] == Y + "/moban/"
    assert Y + "/moban/list-4.html" not in fake.urls()  # page 3 brought nothing new
    short = FakeFetcher(pages)
    assert len(list(slides.ypppt(short, {"sections": ["/moban"], "max_list_pages": 1}))) == 2
    with pytest.raises(ValueError):
        list(slides.ypppt(FakeFetcher({}), {"sections": "templates"}))


def test_catalog_gates_slide_sources_by_terms():
    assert catalog.load("edu_courseware") is slides.webplus and catalog.load("zenodo_cjk") is slides.zenodo_cjk
    assert catalog.load("valse_slides") is slides.valse and catalog.load("ypppt_courseware") is slides.ypppt
    with pytest.raises(PermissionError):
        list(catalog.discover("ypppt_courseware", FakeFetcher(YPAGES)))
    found = list(catalog.discover("ypppt_courseware", FakeFetcher(YPAGES), {"max_items": 1}, accepted={"study_only"}))
    assert found[0]["source"] == "ypppt_courseware" and found[0]["terms"] == "study_only"
    opened = list(catalog.discover("zenodo_cjk", FakeFetcher({slides.ZENODO10K: zenodo_parquet(ZROWS[:2])})))
    assert opened[0]["terms"] == "open"


# --- archives ---

def gbk_zip(members_):
    """A ZIP whose non-ASCII names are stored as GBK bytes without the UTF-8 flag (as Chinese tools do)."""
    buffer, renames = io.BytesIO(), {}
    with zipfile.ZipFile(buffer, "w") as archive:
        for index, (name, data) in enumerate(members_):
            raw = name.encode("gbk")
            stand_in = (f"{index:02d}" + "_" * len(raw))[:len(raw)]
            renames[stand_in.encode()] = raw
            archive.writestr(stand_in if not name.isascii() else name, data)
    data = buffer.getvalue()
    for stand_in, raw in renames.items():
        data = data.replace(stand_in, raw) if stand_in != raw else data
    return data


def rar4(members_):
    """A RAR 4 archive with stored (uncompressed) members and UTF-8 names."""
    def block(kind, flags, body=b"", data=b""):
        header = struct.pack("<BHH", kind, flags, 7 + len(body)) + body
        return struct.pack("<H", zlib.crc32(header) & 0xFFFF) + header + data

    out = b"Rar!\x1a\x07\x00" + block(0x73, 0, b"\0" * 6)
    for name, data in members_:
        raw = name.encode()
        body = struct.pack("<IIBIIBBHI", len(data), len(data), 2, zlib.crc32(data), 40 << 25 | 1 << 21 | 1 << 16, 20,
                           0x30, len(raw), 0x20) + raw
        out += block(0x74, 0x8000 | (0 if name.isascii() else 0x200), body, data)
    return out + block(0x7B, 0x4000)


DECK, PDF = b"PK\x03\x04 deck", b"%PDF-1.4 slides"


def test_zip_members_recover_gbk_names_flatten_paths_and_whitelist(tmp_path):
    archive = gbk_zip([("背影PPT课件8/背影PPT课件8.pptx", DECK), ("../../etc/讲义.pdf", PDF),
                       ("__MACOSX/._背影.pptx", b"junk"), ("说明.txt", b"readme"), ("logo.png", b"\x89PNG")])
    assert members(archive) == [("背影PPT课件8.pptx", DECK), ("讲义.pdf", PDF)]
    path = tmp_path / "a.zip"
    path.write_bytes(archive)
    assert members(str(path)) == members(path) == members(archive)
    utf8 = io.BytesIO()
    with zipfile.ZipFile(utf8, "w") as handle:  # zipfile itself sets the UTF-8 flag
        handle.writestr("第一章.pptx", DECK)
        handle.writestr("dir/第一章.pptx", DECK + b"2")
    assert members(utf8.getvalue()) == [("第一章.pptx", DECK), ("第一章_1.pptx", DECK + b"2")]


def test_member_caps_and_damaged_archives():
    archive = gbk_zip([(f"{i}.pdf", PDF + bytes([i])) for i in range(5)])
    assert len(members(archive, max_members=2)) == 2
    assert members(archive, max_member=len(PDF)) == []
    assert len(members(archive, max_total=2 * len(PDF) + 2)) == 2
    assert members(gbk_zip([("a.pptx", DECK)]), extensions=("pdf",)) == []
    logged = []
    assert members(b"PK\x03\x04" + b"\0" * 40, log=logged.append) == [] and logged
    assert members(b"not an archive at all", log=logged.append) == []
    truncated = gbk_zip([("a.pdf", PDF * 50), ("b.pdf", PDF)])
    assert members(truncated[:-200]) == []  # no central directory: unreadable as a whole


def test_rar_members_through_libarchive():
    pytest.importorskip("ctypes")
    try:
        from data_pipeline.documents.unpack import _library
        _library()
    except OSError:
        pytest.skip("libarchive not installed")
    archive = rar4([("课件/第一章.pptx", DECK), ("../../evil.pdf", PDF), ("notes.txt", b"hello"), ("x/第一章.pptx", DECK)])
    assert members(archive) == [("第一章.pptx", DECK), ("evil.pdf", PDF), ("第一章_1.pptx", DECK)]
    assert members(archive, max_member=len(PDF) - 1) == [("第一章.pptx", DECK), ("第一章_1.pptx", DECK)]


# --- refill ---

def template_deck():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # title only
    slide.shapes.title.text = "单击此处添加标题"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame
    run = box.paragraphs[0].add_run()
    run.text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit"
    run.font.bold, run.font.size = True, Pt(20)
    second = box.add_paragraph()
    second.text = "依法治国是党领导人民治理国家的基本方略。"
    third = box.add_paragraph()
    third.text = "LOGO"
    filler = box.add_paragraph()
    filler.text = "Fully editable shapes and colors for every business presentation"
    fourth = box.add_paragraph()
    fourth.text = "更多模板：www.ypppt.com"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1)).text_frame.text = "请替换文字内容"
    table = slide.shapes.add_table(1, 2, Inches(5), Inches(5), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "输入标题"
    table.cell(0, 1).text = "2025年"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


SENTENCES = ["第一条为了规范行政执法行为，制定本条例。", "国家机关应当依法履行职责，接受社会监督。",
             "本条例自公布之日起施行。", "详见 www.1ppt.com 模板下载。", "Plain English sentence.",
             "各级人民政府应当加强对安全生产工作的领导，支持、督促各有关部门依法履行安全生产监督管理职责。"]


def paragraphs(data):
    from pptx import Presentation
    from data_pipeline.documents.refill import _text_frames
    deck = Presentation(io.BytesIO(data))
    return [p for slide in deck.slides for frame in _text_frames(slide.shapes) for p in frame.paragraphs]


def test_refill_replaces_placeholders_and_filler_keeps_formatting_and_strips_watermarks():
    from data_pipeline.documents.refill import PLACEHOLDER, refill, width
    original = template_deck()
    data, stats = refill(original, SENTENCES, seed=3)
    assert stats == dict(stats, slides=1, placeholder=4, latin=1, watermark=1, replaced=5)
    texts = ["".join(r.text for r in p.runs) for p in paragraphs(data)]
    assert "依法治国是党领导人民治理国家的基本方略。" in texts and "LOGO" in texts and "2025年" in texts
    assert not any(t in texts for t in ("单击此处添加标题", "请替换文字内容", "输入标题"))
    assert "" in texts and not any("ypppt" in t or "1ppt" in t or "English" in t for t in texts)
    lorem = next(p for p in paragraphs(data) if p.runs and p.runs[0].font.bold)
    assert lorem.runs[0].font.size.pt == 20 and abs(len(lorem.runs[0].text) - 28) <= 6
    title = texts[0]
    assert abs(width(title) - 8) <= 2 and all("一" <= c <= "鿿" or c in "，。、" for c in title)
    assert stats["cjk_after"] > stats["cjk_before"]
    assert refill(original, SENTENCES, seed=3)[1] == stats
    assert all(PLACEHOLDER.search(t) for t in ("输入你需要的文本内容", "文本标题", "ADD YOUR LOGO HERE", "202X年工作总结"))
    assert not any(PLACEHOLDER.search(t) for t in ("2020年工作总结", "古诗起源", "PART ONE", "实验验证"))
    again = refill(original, SENTENCES, 3)[0]
    assert texts == ["".join(r.text for r in p.runs) for p in paragraphs(again)]
    kept, latin_off = refill(original, SENTENCES, seed=3, latin=False)
    assert latin_off["latin"] == 0 and latin_off["placeholder"] == 4
    with pytest.raises(ValueError):
        refill(original, ["see www.example.com", "English only"], seed=0)


def test_refill_corpus_takes_similar_lengths():
    import random
    from data_pipeline.documents.refill import Corpus, latin_filler
    corpus = Corpus(SENTENCES, random.Random(0))
    assert all("www" not in item for item in corpus.items)
    assert all(n * 0.75 <= len(corpus.take(n)) <= n * 1.25 + 1 for n in (2, 4, 9, 11, 16, 40, 120))
    assert len(corpus.take(300)) == 300
    assert latin_filler("Your content to play here, or through your copy") and not latin_filler("PART 01")
    assert not latin_filler("2023年度 Annual Report of the Company")
